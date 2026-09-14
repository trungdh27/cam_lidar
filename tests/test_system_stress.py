import asyncio
import hashlib
import json
import os
import sys
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QProcess, Qt
from PySide6.QtWidgets import QApplication, QFrame

from desktop_app.stress.catalog import DEFAULT_VD_SOURCE, StressCatalog, StressCatalogError
from desktop_app.stress.collectors import EvidenceManager
from desktop_app.stress.evidence_plan import build_evidence_plan
from desktop_app.stress.log_tail import IncrementalLogTail
from desktop_app.stress.models import (
    CollectorStatus,
    EnvironmentStatus,
    EvidenceDefinition,
    ExecutionType,
    PreTestCheck,
    PreTestStatus,
    RuntimeStatus,
    StressTestDefinition,
)
from desktop_app.stress.pretest import (
    BASELINE_DEPENDENCIES,
    PreTestEnvironmentRunner,
    ReadinessResult,
    dependencies_for_tests,
    evaluate_readiness,
)
from desktop_app.stress.runner import StressTestRunner, scan_history
from desktop_app.stress.session import StressSessionManager
from desktop_app.ui.main_window import MainWindow
from desktop_app.ui.system_stress_page import StatusBadge, StressEvidenceViewer, SystemStressPage


SOURCE_HASH = hashlib.sha256(DEFAULT_VD_SOURCE.read_bytes()).hexdigest()


def process_until(predicate, timeout=5.0):
    app = QApplication.instance()
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return True
        time.sleep(0.01)
    app.processEvents()
    return bool(predicate())


def definition(test_id="ST-DUMMY-001", *, evidence=None, program=None, arguments=(), target_cpu_percent=None):
    return StressTestDefinition(
        test_id=test_id,
        group="Dummy",
        test_name=f"Dummy {test_id}",
        source_fields={"Test ID": test_id, "Test Name": f"Dummy {test_id}"},
        duration_text="1 sec",
        duration_seconds=1,
        execution_type=ExecutionType.AUTO if program else ExecutionType.GUIDED,
        evidence_plan=tuple(evidence or build_evidence_plan("CPU/GPU/AI Stress", "CPU")),
        workload_program=program,
        workload_arguments=tuple(arguments),
        workload_description="Synthetic test command" if program else None,
        target_cpu_percent=target_cpu_percent,
    )


class CatalogTests(unittest.TestCase):
    def test_05_06_vd_catalog_loads_exactly_100_unique_cases(self):
        catalog = StressCatalog.load_vd()
        self.assertEqual(len(catalog.definitions), 100)
        self.assertEqual(len({item.test_id for item in catalog.definitions}), 100)
        self.assertTrue(all(item.test_id.startswith("ST-") for item in catalog.definitions))
        self.assertEqual(catalog.encoding, "xlsx-unicode")

    def test_xlsx_preserves_vietnamese_unicode_exactly(self):
        catalog = StressCatalog.load_vd()
        source = catalog.definitions[0].source_fields
        self.assertEqual(catalog.encoding, "xlsx-unicode")
        self.assertIn("Nhóm", source)
        self.assertEqual(
            source["Purpose"],
            "Kiểm tra robot có hoạt động ổn định khi CPU phải xử lý tải khoảng 30%.",
        )
        self.assertIn("Robot bật bình thường", source["Pre-condition"])
        self.assertIn("pin ≥50%", source["Pre-condition"])
        self.assertIn("vị trí an toàn", source["Pre-condition"])
        self.assertIn("Thời gian", source)
        self.assertIn("Giá trị đo thực tế", source)

    def test_authoritative_xlsx_has_no_legacy_encoding_warning(self):
        catalog = StressCatalog.load_vd()
        self.assertEqual(catalog.source_warnings, [])
        self.assertEqual(catalog.source_path.suffix, ".xlsx")

    def test_07_duplicate_ids_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "duplicate.xlsx"
            workbook = openpyxl.load_workbook(DEFAULT_VD_SOURCE)
            sheet = workbook["Stress system"]
            sheet["A3"] = sheet["A2"].value
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(StressCatalogError, "Duplicate Test ID"):
                StressCatalog.load_vd(path)

    def test_08_invalid_xlsx_row_is_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Test ID", "Group", "Test Name"])
            sheet.append(["ST-X-001", "Only group", None])
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(StressCatalogError, "Invalid XLSX row 2"):
                StressCatalog.load_vd(path, expected_count=1)

    def test_required_column_error_is_clear(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "missing.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["id", "name"])
            sheet.append(["1", "x"])
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(StressCatalogError, "exactly one worksheet"):
                StressCatalog.load_vd(path, expected_count=1)

    def test_legacy_csv_is_not_accepted_as_authoritative(self):
        legacy = DEFAULT_VD_SOURCE.with_suffix(".csv")
        with self.assertRaisesRegex(StressCatalogError, "must be an XLSX workbook"):
            StressCatalog.load_vd(legacy)

    def test_09_10_fresh_status_ignores_source_history(self):
        catalog = StressCatalog.load_vd()
        self.assertTrue(all(item.runtime_status == RuntimeStatus.NOT_RUN for item in catalog.definitions))
        historical = [item for item in catalog.definitions if item.source_fields.get("Pass/Fail")]
        for item in historical:
            self.assertEqual(item.runtime_status, RuntimeStatus.NOT_RUN)

    def test_11_search_works(self):
        catalog = StressCatalog.load_vd()
        result = catalog.filter(search="ST-CPU-003")
        self.assertEqual([item.test_id for item in result], ["ST-CPU-003"])

    def test_12_group_filter_works(self):
        catalog = StressCatalog.load_vd()
        group = catalog.definitions[0].group
        result = catalog.filter(group=group)
        self.assertEqual(len(result), 10)
        self.assertTrue(all(item.group == group for item in result))

    def test_13_status_filter_works(self):
        catalog = StressCatalog.load_vd()
        catalog.definitions[0].runtime_status = RuntimeStatus.RUNNING
        self.assertEqual(catalog.filter(status="RUNNING"), [catalog.definitions[0]])

    def test_14_severity_filter_works(self):
        catalog = StressCatalog.load_vd()
        severity = catalog.definitions[0].severity
        self.assertTrue(severity)
        self.assertTrue(all(item.severity == severity for item in catalog.filter(severity=severity)))

    def test_source_file_remains_unchanged_during_catalog_tests(self):
        self.assertEqual(hashlib.sha256(DEFAULT_VD_SOURCE.read_bytes()).hexdigest(), SOURCE_HASH)


class SessionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.manager = StressSessionManager(self.root)
        self.item = definition()

    def tearDown(self):
        self.temp.cleanup()

    def test_17_18_evidence_root_custom_and_default_models(self):
        self.assertEqual(self.manager.set_evidence_root(self.root), self.root.resolve())
        default = StressSessionManager()
        self.assertEqual(default.evidence_root, Path.home() / "Stress_Test_Logs")

    def test_19_20_session_id_and_folder_creation(self):
        path = self.manager.create_session("VD", now=datetime(2026, 9, 9, 10, 11, 12))
        self.assertEqual(path.name, "VD_20260909_101112")
        self.assertTrue((path / "pre_test_environment").is_dir())
        self.assertTrue((path / "tests").is_dir())

    def test_session_collision_never_overwrites(self):
        first = self.manager.create_session("VD", now=datetime(2026, 9, 9, 10, 11, 12))
        second_manager = StressSessionManager(self.root)
        second = second_manager.create_session("VD", now=datetime(2026, 9, 9, 10, 11, 12))
        self.assertNotEqual(first, second)
        self.assertEqual(second.name, "VD_20260909_101112_02")

    def test_21_22_23_24_attempts_increment_without_overwrite(self):
        self.manager.create_session("VD")
        first = self.manager.create_attempt(self.item)
        marker = first.attempt_dir / "do_not_overwrite.txt"
        marker.write_text("preserve", encoding="utf-8")
        second = self.manager.create_attempt(self.item)
        self.assertEqual(first.attempt, 1)
        self.assertEqual(second.attempt, 2)
        self.assertTrue(marker.is_file())
        self.assertEqual(marker.read_text(), "preserve")

    def test_25_26_27_metadata_files_are_created(self):
        session = self.manager.create_session("VD")
        paths = self.manager.create_attempt(self.item)
        self.assertTrue((session / "session.json").is_file())
        self.assertTrue((paths.attempt_dir / "test_info.json").is_file())
        self.assertTrue((paths.attempt_dir / "evidence_manifest.json").is_file())
        self.assertTrue((paths.artifacts_dir / "video").is_dir())

    def test_28_manifest_status_update(self):
        self.manager.create_session("VD")
        paths = self.manager.create_attempt(self.item)
        manager = EvidenceManager(self.item, paths.attempt_dir)
        manager.records[0].status = CollectorStatus.RUNNING
        self.manager.update_manifest(paths, manager.records, self.item.test_id)
        payload = json.loads((paths.attempt_dir / "evidence_manifest.json").read_text())
        self.assertEqual(payload["evidence"][0]["status"], "RUNNING")

    def test_44_override_is_persisted_with_reasons(self):
        self.manager.create_session("VD")
        self.manager.update_environment(EnvironmentStatus.NOT_READY, ["LiDAR missing"], override=True)
        payload = json.loads((self.manager.session_dir / "session.json").read_text())
        self.assertTrue(payload["environment_override"])
        self.assertEqual(payload["environment_blocking_reasons"], ["LiDAR missing"])

    def test_history_reads_attempt_folders_without_database(self):
        self.manager.create_session("VD")
        paths = self.manager.create_attempt(self.item)
        self.manager.write_result(paths, {"test_id": self.item.test_id, "attempt": 1, "status": "NEEDS_REVIEW"})
        history = scan_history(self.root, self.item.test_id)
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["status"], "NEEDS_REVIEW")


class ReadinessTests(unittest.TestCase):
    def test_ros2_passes_after_sourcing_humble_on_shared_ssh(self):
        class Result:
            def __init__(self, stdout="", stderr="", exit_status=0):
                self.stdout = stdout
                self.stderr = stderr
                self.exit_status = exit_status

        class SharedSSH:
            def __init__(self):
                self.commands = []

            async def run(self, command, timeout):
                self.commands.append((command, timeout))
                if "/opt/ros/humble/setup.bash" in command:
                    return Result("[PASS] ros2\n/rosout\n")
                if "for c in" in command:
                    return Result("[PASS] stress-ng\n[PASS] journalctl\n")
                return Result()

        ssh = SharedSSH()
        payload = asyncio.run(
            PreTestEnvironmentRunner.remote_operation(
                ssh, set(BASELINE_DEPENDENCIES) | {"ros2"}
            )
        )
        ros2_command = next(command for command, _timeout in ssh.commands if "/opt/ros/humble/setup.bash" in command)
        self.assertIn("bash -lc", ros2_command)
        self.assertIn("source /opt/ros/humble/setup.bash", ros2_command)
        self.assertIn("command -v ros2", ros2_command)
        self.assertIn("ros2 --help", ros2_command)
        with tempfile.TemporaryDirectory() as directory:
            result = PreTestEnvironmentRunner.result_from_remote(
                payload, Path(directory), set(BASELINE_DEPENDENCIES)
            )
        ros2 = next(check for check in result.checks if check.id == "tool_ros2")
        self.assertEqual(ros2.status, PreTestStatus.PASS)
        self.assertIn("after sourcing", ros2.actual)

    def test_warning_is_non_blocking_even_for_required_dependency(self):
        checks = [
            PreTestCheck(
                "kernel",
                "Kernel/system errors",
                "Kernel",
                "Review captured warnings",
                PreTestStatus.WARNING,
                "logging",
                True,
            )
        ]
        status, reasons = evaluate_readiness(checks, {"logging"})
        self.assertEqual(status, EnvironmentStatus.READY)
        self.assertEqual(reasons, [])

    def test_required_missing_dependency_blocks(self):
        checks = [
            PreTestCheck(
                "tool_stress-ng",
                "Tool: stress-ng",
                "Tools",
                "Not found in PATH",
                PreTestStatus.MISSING,
                "workload",
            )
        ]
        status, reasons = evaluate_readiness(checks, {"workload"})
        self.assertEqual(status, EnvironmentStatus.NOT_READY)
        self.assertEqual(reasons, ["Tool: stress-ng: MISSING (Not found in PATH)"])

    def test_unrelated_not_configured_dependency_does_not_block(self):
        checks = [
            PreTestCheck(
                "camera",
                "Camera",
                "Target",
                "Required stream/topic not configured",
                PreTestStatus.NOT_CONFIGURED,
                "camera",
            )
        ]
        status, reasons = evaluate_readiness(checks, {"workload"})
        self.assertEqual(status, EnvironmentStatus.READY)
        self.assertEqual(reasons, [])

    def test_selected_cpu_readiness_uses_only_its_dependencies(self):
        cpu_test = StressCatalog.load_vd().get("ST-CPU-001")
        dependencies = set(BASELINE_DEPENDENCIES) | dependencies_for_tests([cpu_test])
        self.assertEqual(dependencies, {"connection", "logging", "workload"})
        checks = [
            PreTestCheck("connection", "DUT connection", "Target", "Connected", PreTestStatus.PASS, "connection"),
            PreTestCheck("logging", "Logging", "Resources", "Writable", PreTestStatus.PASS, "logging"),
            PreTestCheck("stress-ng", "Tool: stress-ng", "Tools", "Available", PreTestStatus.PASS, "workload"),
            PreTestCheck("camera", "Camera", "Target", "Not configured", PreTestStatus.NOT_CONFIGURED, "camera"),
            PreTestCheck("lidar", "LiDAR", "Target", "Not configured", PreTestStatus.NOT_CONFIGURED, "lidar"),
            PreTestCheck("estop", "E-stop", "Target", "Not configured", PreTestStatus.NOT_CONFIGURED, "estop"),
        ]
        status, reasons = evaluate_readiness(checks, dependencies)
        self.assertEqual(status, EnvironmentStatus.READY)
        self.assertEqual(reasons, [])

    def test_42_ready_case_only_considers_relevant_dependencies(self):
        checks = [
            PreTestCheck("logging", "Logging", "Resource", "Writable", PreTestStatus.PASS, "logging", True),
            PreTestCheck("can", "CAN", "Target", "Not configured", PreTestStatus.NOT_CONFIGURED, "can"),
        ]
        status, reasons = evaluate_readiness(checks, {"logging"})
        self.assertEqual(status, EnvironmentStatus.READY)
        self.assertEqual(reasons, [])

    def test_43_not_ready_case_reports_exact_reason(self):
        checks = [PreTestCheck("can", "CAN", "Target", "can0 absent", PreTestStatus.FAIL, "can")]
        status, reasons = evaluate_readiness(checks, {"can"})
        self.assertEqual(status, EnvironmentStatus.NOT_READY)
        self.assertIn("CAN: FAIL (can0 absent)", reasons)

    def test_45_missing_tool_is_reported(self):
        with tempfile.TemporaryDirectory() as directory, patch("desktop_app.stress.pretest.shutil.which", return_value=None):
            runner = PreTestEnvironmentRunner(command_runner=lambda _cmd: type("R", (), {"stdout": "", "stderr": "", "returncode": 0})(), baseline_duration_sec=0.002, baseline_interval_sec=0.001)
            result = runner.run_local(Path(directory), {"workload", "logging"})
            stress = next(check for check in result.checks if check.id == "tool_stress-ng")
            self.assertEqual(stress.status, PreTestStatus.MISSING)
            self.assertEqual(result.status, EnvironmentStatus.NOT_READY)

    def test_46_unconfigured_dependency_is_not_pass(self):
        with tempfile.TemporaryDirectory() as directory:
            result = PreTestEnvironmentRunner(baseline_duration_sec=0.002, baseline_interval_sec=0.001).run_local(Path(directory), {"lidar", "logging"})
            lidar = next(check for check in result.checks if check.id == "lidar")
            self.assertEqual(lidar.status, PreTestStatus.NOT_CONFIGURED)
            self.assertNotEqual(lidar.status, PreTestStatus.PASS)
            self.assertEqual(result.status, EnvironmentStatus.NOT_READY)

    def test_pretest_writes_all_required_baseline_files(self):
        with tempfile.TemporaryDirectory() as directory:
            result = PreTestEnvironmentRunner(baseline_duration_sec=0.002, baseline_interval_sec=0.001).run_local(Path(directory), {"logging"})
            self.assertTrue((Path(directory) / "readiness.json").is_file())
            for index in range(1, 9):
                self.assertEqual(len(list(Path(directory).glob(f"{index:02d}_*.log"))), 1)
            self.assertTrue((Path(directory) / "baseline.json").is_file())
            self.assertTrue((Path(directory) / "baseline_summary.log").is_file())
            for name in (
                "09_system_metrics_baseline.log",
                "10_tegrastats_baseline.log",
                "11_kernel_baseline.log",
                "12_journal_baseline.log",
            ):
                self.assertTrue((Path(directory) / name).is_file())


class LogTailTests(unittest.TestCase):
    def test_32_incremental_tail_reads_only_new_bytes(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "log.txt"
            path.write_text("old\n", encoding="utf-8")
            tail = IncrementalLogTail()
            self.assertEqual(tail.select(path), "old\n")
            with path.open("a") as stream:
                stream.write("new\n")
            self.assertEqual(tail.read_new(), "new\n")
            self.assertEqual(tail.read_new(), "")

    def test_33_truncation_and_recreation_are_handled(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "log.txt"
            path.write_text("long original content\n", encoding="utf-8")
            tail = IncrementalLogTail()
            tail.select(path)
            path.write_text("short\n", encoding="utf-8")
            self.assertIn("truncated or replaced", tail.read_new())
            path.unlink()
            path.write_text("rotated\n", encoding="utf-8")
            self.assertIn("rotated", tail.read_new())

    def test_initial_load_is_bounded_to_requested_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "log.txt"
            path.write_text("".join(f"{i}\n" for i in range(1000)), encoding="utf-8")
            content = IncrementalLogTail().select(path, initial_lines=10)
            self.assertEqual(len(content.splitlines()), 10)
            self.assertTrue(content.startswith("990\n"))


class QtStressTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_01_02_03_04_page_constructs_with_exact_tabs_and_pretest_first(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            self.assertEqual(page.tabs.count(), 4)
            self.assertEqual([page.tabs.tabText(i) for i in range(4)], list(SystemStressPage.TAB_NAMES))
            self.assertIs(page.tabs.widget(0), page.pretest_page)
            page.close()

    def test_compact_session_header_and_catalog_details_default_state(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            cards = [frame for frame in page.findChildren(QFrame) if frame.objectName() == "StressPanel"]
            self.assertIn(108, [card.maximumHeight() for card in cards])
            self.assertEqual(page.vd_page.details_view.toPlainText(), "")
            self.assertEqual(
                page.vd_page.details_view.placeholderText(),
                "Select a test case to view details.",
            )
            self.assertEqual(page.vd_page.details_stack.currentIndex(), 0)
            self.assertEqual(page.vd_page.selected_label.text(), "Selected 0  •  Visible 100  •  Total 100")
            self.assertEqual(page.vd_page.message_label.text(), "")
            page.close()

    def test_pretest_uses_compact_split_detail_area(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            self.assertTrue(page.pretest_page.dut_label.isHidden())
            self.assertIn("Select a check", page.pretest_page.detail_area.toPlainText())
            page.pretest_page.view_details_button.click()
            self.assertEqual(page.pretest_page.pretest_detail_tabs.currentIndex(), 1)
            self.assertIn("Run the baseline checks", page.pretest_page.dut_info_view.toPlainText())
            page.close()

    def test_15_16_details_and_dynamic_evidence_plan_render(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            page.vd_page.table.selectRow(0)
            page.vd_page.render_details()
            self.assertIn(page.catalog.definitions[0].test_id, page.vd_page.details_view.toPlainText())
            self.assertEqual(
                [page.vd_page.detail_tabs.tabText(i) for i in range(4)],
                ["OVERVIEW", "PROCEDURE", "CRITERIA", "EVIDENCE"],
            )
            self.assertIn("Kiểm tra", page.vd_page.overview_view.toPlainText())
            self.assertIn("1.", page.vd_page.procedure_view.toPlainText())
            expected = len(page.catalog.definitions[0].evidence_plan)
            self.assertEqual(page.vd_page.evidence_plan_list.count(), expected)
            page.close()

    def test_status_badges_use_readable_labels_and_semantic_tones(self):
        badge = StatusBadge(RuntimeStatus.NOT_RUN)
        self.assertEqual(badge.text(), "NOT RUN")
        self.assertEqual(badge.property("tone"), "neutral")
        badge.set_status(EnvironmentStatus.READY)
        self.assertEqual(badge.property("tone"), "success")
        badge.set_status(EnvironmentStatus.NOT_READY)
        self.assertEqual(badge.property("tone"), "error")

    def test_pretest_labels_baseline_and_selected_test_readiness(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            result = ReadinessResult(
                EnvironmentStatus.READY,
                [
                    PreTestCheck("connection", "DUT connection", "Target", "Connected", PreTestStatus.PASS, "connection"),
                    PreTestCheck("logging", "Logging", "Resources", "Writable", PreTestStatus.PASS, "logging"),
                    PreTestCheck("stress-ng", "Tool: stress-ng", "Tools", "Available", PreTestStatus.PASS, "workload"),
                    PreTestCheck("camera", "Camera", "Target", "Not configured", PreTestStatus.NOT_CONFIGURED, "camera"),
                ],
                [],
                {},
                "now",
            )
            page.pretest_page.result = result
            page.pretest_page._render(result)
            self.assertEqual(page.pretest_page.readiness_label.text(), "BASELINE READY")
            page.pretest_page.set_required_tests([page.catalog.get("ST-CPU-001")])
            self.assertEqual(page.pretest_page.readiness_label.text(), "READY FOR SELECTED TEST")
            self.assertEqual(page.pretest_page.blocking_reasons, [])
            page.close()

    def test_compact_filter_bar_has_only_three_filters_and_reset(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            self.assertFalse(hasattr(page.vd_page, "search_edit"))
            self.assertFalse(hasattr(page.vd_page, "quick_buttons"))
            group = page.catalog.definitions[0].group
            page.vd_page.group_combo.setCurrentText(group)
            self.assertEqual(page.vd_page.table.rowCount(), 10)
            page.vd_page.status_combo.setCurrentText("NOT_RUN")
            page.vd_page.severity_combo.setCurrentText("Major")
            self.assertEqual(page.vd_page.table.rowCount(), 10)
            page.vd_page.reset_button.click()
            self.assertEqual(page.vd_page.group_combo.currentText(), "All Groups")
            self.assertEqual(page.vd_page.status_combo.currentText(), "All")
            self.assertEqual(page.vd_page.severity_combo.currentText(), "All")
            self.assertEqual(page.vd_page.table.rowCount(), 100)
            page.close()

    def test_sort_columns_visibility_refresh_and_footer(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            source_order = [item.test_id for item in page.catalog.definitions]
            page.vd_page.table.sortItems(1, Qt.SortOrder.DescendingOrder)
            self.assertNotEqual(page.vd_page.table.item(0, 1).text(), source_order[0])
            self.assertEqual([item.test_id for item in page.catalog.definitions], source_order)
            page.vd_page.column_actions[2].setChecked(False)
            self.assertTrue(page.vd_page.table.isColumnHidden(2))
            page.vd_page.column_actions[2].setChecked(True)
            self.assertFalse(page.vd_page.table.isColumnHidden(2))
            page.catalog.definitions[0].runtime_status = RuntimeStatus.RUNNING
            page.vd_page.refresh_button.click()
            self.assertEqual(page.catalog.definitions[0].runtime_status, RuntimeStatus.RUNNING)
            self.assertIn("Total 100", page.vd_page.selected_label.text())
            page.close()

    def test_pretest_failure_is_actionable_and_table_height_adapts(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            result = ReadinessResult(
                EnvironmentStatus.NOT_READY,
                [PreTestCheck("target_connection", "DUT connection", "Target", "Not connected", PreTestStatus.FAIL, "logging", True)],
                ["DUT connection: FAIL (Not connected)"],
                {"connection": "not connected"},
                "now",
            )
            page.session_manager.create_session("VD")
            page.pretest_page._completed(result)
            page.pretest_page.check_table.selectRow(0)
            self.assertLess(page.pretest_page.check_table.maximumHeight(), 100)
            self.assertFalse(page.pretest_page.go_dashboard_button.isHidden())
            self.assertFalse(page.pretest_page.override_button.isHidden())
            self.assertIn("Dashboard", page.pretest_page.detail_area.toPlainText())
            page.close()

    def test_filters_update_vd_table(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            definition = page.catalog.definitions[10]
            page.vd_page.group_combo.setCurrentText(definition.group)
            expected = sum(item.group == definition.group for item in page.catalog.definitions)
            self.assertEqual(page.vd_page.table.rowCount(), expected)
            definition.runtime_status = RuntimeStatus.RUNNING
            page.vd_page.status_combo.setCurrentText("RUNNING")
            self.assertEqual(page.vd_page.table.rowCount(), 1)
            page.close()

    def test_29_variable_evidence_list_counts(self):
        with tempfile.TemporaryDirectory() as directory:
            viewer = StressEvidenceViewer()
            item = definition(evidence=(EvidenceDefinition("a", "A", "log", "manual", "A", manual_required=True),))
            manager = EvidenceManager(item, Path(directory))
            viewer.set_evidence(manager.records, directory)
            self.assertEqual(viewer.evidence_list.count(), 1)
            item2 = definition(evidence=tuple(EvidenceDefinition(str(i), f"Evidence {i}", "log", "manual", "manual", manual_required=True) for i in range(10)))
            manager2 = EvidenceManager(item2, Path(directory))
            viewer.set_evidence(manager2.records, directory)
            self.assertEqual(viewer.evidence_list.count(), 10)
            viewer.close()

    def test_30_31_34_35_36_viewer_selection_timer_bounds_refresh_and_autoscroll(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            evidence = tuple(EvidenceDefinition(name, name.upper(), "log", "manual", name, manual_required=True) for name in ("one", "two"))
            item = definition(evidence=evidence)
            manager = EvidenceManager(item, root)
            for record in manager.records:
                path = root / record.file
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(record.id + "\n", encoding="utf-8")
            viewer = StressEvidenceViewer()
            viewer.set_evidence(manager.records, root)
            self.assertEqual(viewer.refresh_timer.interval(), 10000)
            self.assertEqual(viewer.log_view.document().maximumBlockCount(), 2000)
            viewer.evidence_list.setCurrentRow(1)
            self.assertIn("two", viewer.log_view.toPlainText())
            with (root / manager.records[1].file).open("a") as stream:
                stream.write("appended\n")
            viewer.refresh_button.click()
            self.assertIn("appended", viewer.log_view.toPlainText())
            viewer.auto_scroll.setChecked(False)
            self.assertFalse(viewer.auto_scroll.isChecked())
            viewer.close()

    def test_system_metrics_summary_raw_switching_and_incremental_refresh(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            evidence = (
                EvidenceDefinition("system_metrics", "System Metrics", "log", "system_metrics", "metrics", 5),
            )
            item = definition(
                "ST-CPU-001",
                evidence=evidence,
                program="stress-ng",
                arguments=("--cpu", "0", "--timeout", "1s"),
                target_cpu_percent=30,
            )
            manager = EvidenceManager(item, root)
            record = manager.records[0]
            record.status = CollectorStatus.RUNNING
            path = root / record.file
            path.parent.mkdir(parents=True, exist_ok=True)
            first = (
                "[stdout] 2026-09-10T10:32:54+00:00\n"
                "[stdout] cpu 100 0 50 800 50 0 0 0\n"
                "[stdout] cpu0 50 0 25 400 25 0 0 0\n"
                "[stdout] 10:32:54 up 1:19, 1 user, load average: 17.46, 16.30, 14.76\n"
                "[stdout] Mem: 65893445632 6302867456 1000 0 0 57575608320\n"
                "[stdout] Swap: 32942342144 524288 32941817856\n"
            )
            second = (
                "[stdout] 2026-09-10T10:32:59+00:00\n"
                "[stdout] cpu 140 0 60 840 60 0 0 0\n"
                "[stdout] cpu0 70 0 30 420 30 0 0 0\n"
                "[stdout] 10:32:59 up 1:19, 1 user, load average: 17.46, 16.30, 14.76\n"
                "[stdout] Mem: 65893445632 6302867456 1000 0 0 57575608320\n"
                "[stdout] Swap: 32942342144 524288 32941817856\n"
            )
            path.write_text(first, encoding="utf-8")
            original = path.read_bytes()
            viewer = StressEvidenceViewer()
            viewer.set_evidence(manager.records, root, item)
            self.assertEqual(viewer.refresh_timer.interval(), 10000)
            self.assertEqual(viewer.content_stack.currentIndex(), 0)
            self.assertIn("SYSTEM METRICS", viewer.summary_view.toPlainText())
            self.assertIn("Target CPU     30 %", viewer.summary_view.toPlainText())
            self.assertIn("Waiting for next sample", viewer.summary_view.toPlainText())
            self.assertNotIn("\n", viewer.evidence_list.item(0).text())
            self.assertIn("RUNNING", viewer.evidence_list.item(0).text())

            viewer.raw_button.click()
            self.assertEqual(viewer.content_stack.currentIndex(), 1)
            self.assertEqual(viewer.log_view.toPlainText(), first)
            with path.open("a", encoding="utf-8") as stream:
                stream.write(second)
            viewer.refresh_now()
            self.assertIn(second.rstrip(), viewer.log_view.toPlainText())
            viewer.summary_button.click()
            self.assertEqual(viewer.content_stack.currentIndex(), 0)
            self.assertIn("Current CPU    50.0 %", viewer.summary_view.toPlainText())
            self.assertEqual(path.read_bytes(), original + second.encode("utf-8"))
            viewer.close()

    def test_manual_evidence_and_kernel_warning_have_useful_summaries(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            manual_item = definition(
                evidence=(EvidenceDefinition("manual", "CPU / GPU / AI Measurements", "log", "manual", "manual", manual_required=True),)
            )
            manual_manager = EvidenceManager(manual_item, root)
            manual_path = root / manual_manager.records[0].file
            manual_path.parent.mkdir(parents=True, exist_ok=True)
            manual_path.touch()
            viewer = StressEvidenceViewer()
            viewer.set_evidence(manual_manager.records, root, manual_item)
            self.assertIn("MANUAL EVIDENCE REQUIRED", viewer.summary_view.toPlainText())
            self.assertIn("not collected automatically", viewer.summary_view.toPlainText())
            self.assertFalse(viewer.open_artifact_folder_button.isHidden())

            kernel_item = definition(
                evidence=(EvidenceDefinition("kernel", "Kernel Error Log", "log", "dmesg", "kernel"),)
            )
            kernel_manager = EvidenceManager(kernel_item, root)
            kernel_record = kernel_manager.records[0]
            kernel_record.status = CollectorStatus.WARNING
            kernel_path = root / kernel_record.file
            kernel_path.write_text(
                "[stderr] 2026-09-10T10:32:54+00:00 thermal warning\n[stderr] driver warning\n",
                encoding="utf-8",
            )
            viewer.set_evidence(kernel_manager.records, root, kernel_item)
            summary = viewer.summary_view.toPlainText()
            self.assertIn("KERNEL ERROR LOG", summary)
            self.assertIn("Warning Count   2", summary)
            self.assertIn("thermal warning", summary)
            viewer.close()

    def test_back_to_catalog_preserves_running_execution_and_attempt(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            item = definition(
                "ST-SAFE-001",
                evidence=(EvidenceDefinition("system_metrics", "System Metrics", "log", "system_metrics", "metrics", 0.05),),
            )
            self.assertTrue(page.runner.enqueue([item], environment_status=EnvironmentStatus.READY))
            self.assertTrue(page.runner.start_current())
            self.assertTrue(process_until(lambda: page.runner.evidence_manager.active_count() == 1))
            manager = page.runner.evidence_manager
            paths = page.runner.current_paths
            attempt_dirs = list((paths.session_dir / "tests" / item.test_id).glob("attempt_*"))
            queue_before = list(page.runner.queue)
            elapsed_before = page.runner.elapsed_seconds()

            page.vd_page.back_to_list_button.click()
            self.assertEqual(page.vd_page.view_stack.currentIndex(), 0)
            self.assertEqual(item.runtime_status, RuntimeStatus.RUNNING)
            self.assertIs(page.runner.evidence_manager, manager)
            self.assertEqual(manager.active_count(), 1)
            self.assertIs(page.runner.current_paths, paths)
            self.assertFalse(page.vd_page.active_banner.isHidden())
            self.assertIn(item.test_id, page.vd_page.active_test_label.text())
            self.assertFalse(page.vd_page.run_button.isEnabled())
            self.assertEqual(page.vd_page.run_button.toolTip(), "A stress test is already running.")
            self.assertTrue(process_until(lambda: page.runner.elapsed_seconds() > elapsed_before, timeout=2.5))

            page.run_selected([page.catalog.definitions[0]])
            self.assertEqual(page.vd_page.message_label.text(), "A stress test is already running.")
            self.assertEqual(page.runner.queue, queue_before)
            self.assertIs(page.runner.current_paths, paths)
            self.assertEqual(
                list((paths.session_dir / "tests" / item.test_id).glob("attempt_*")),
                attempt_dirs,
            )
            page.vd_page.view_running_button.click()
            self.assertEqual(page.vd_page.view_stack.currentIndex(), 2)
            self.assertIs(page.runner.evidence_manager, manager)
            self.assertIs(page.runner.current_paths, paths)

            page.runner.stop()
            self.assertTrue(process_until(lambda: item.runtime_status == RuntimeStatus.STOPPED))
            page.close()

    def test_37_dummy_collectors_run_concurrently(self):
        with tempfile.TemporaryDirectory() as directory:
            evidence = (
                EvidenceDefinition("one", "One", "log", "system_metrics", "synthetic", 0.05),
                EvidenceDefinition("two", "Two", "log", "system_metrics", "synthetic", 0.05),
            )
            manager = EvidenceManager(definition(evidence=evidence), Path(directory))
            manager.start_all()
            self.assertTrue(process_until(lambda: manager.active_count() == 2))
            manager.stop_all()
            self.assertTrue(process_until(lambda: manager.active_count() == 0))

    def test_38_collector_failure_does_not_crash_other_collector(self):
        with tempfile.TemporaryDirectory() as directory:
            evidence = (
                EvidenceDefinition("bad", "Bad", "log", "unknown_collector", "bad"),
                EvidenceDefinition("good", "Good", "log", "system_metrics", "good", 0.05),
            )
            manager = EvidenceManager(definition(evidence=evidence), Path(directory))
            manager.start_all()
            self.assertTrue(process_until(lambda: manager.collectors["bad"].record.status in {CollectorStatus.COMPLETED, CollectorStatus.ERROR}))
            self.assertTrue(manager.collectors["good"].running)
            manager.stop_all()
            self.assertTrue(process_until(lambda: not manager.collectors["good"].running))

    def test_39_40_stop_terminates_owned_process_without_zombie(self):
        with tempfile.TemporaryDirectory() as directory:
            evidence = (EvidenceDefinition("metrics", "Metrics", "log", "system_metrics", "metrics", 0.05),)
            manager = EvidenceManager(definition(evidence=evidence), Path(directory))
            collector = manager.collectors["metrics"]
            manager.start_all()
            self.assertTrue(process_until(lambda: collector.running))
            manager.stop_all()
            self.assertTrue(process_until(lambda: not collector.running))
            self.assertEqual(collector.process.state(), QProcess.ProcessState.NotRunning)

    def test_41_queue_is_sequential_and_requires_each_start(self):
        with tempfile.TemporaryDirectory() as directory:
            manager = StressSessionManager(directory)
            runner = StressTestRunner(manager)
            evidence = (EvidenceDefinition("workload", "Workload", "log", "workload", "synthetic"),)
            first = definition("ST-DUMMY-001", evidence=evidence, program=sys.executable, arguments=("-c", "print('one')"))
            second = definition("ST-DUMMY-002", evidence=evidence, program=sys.executable, arguments=("-c", "print('two')"))
            confirmations = []
            runner.confirmation_required.connect(lambda item, _paths: confirmations.append(item.test_id))
            self.assertTrue(runner.enqueue([first, second], environment_status=EnvironmentStatus.READY))
            self.assertEqual(confirmations, [first.test_id])
            self.assertEqual(second.runtime_status, RuntimeStatus.WAITING)
            runner.start_current()
            self.assertTrue(process_until(lambda: len(confirmations) == 2))
            self.assertEqual(first.runtime_status, RuntimeStatus.NEEDS_REVIEW)
            self.assertEqual(second.runtime_status, RuntimeStatus.STARTING)
            self.assertIsNone(runner.evidence_manager)

    def test_runner_stop_writes_stopped_result(self):
        with tempfile.TemporaryDirectory() as directory:
            session = StressSessionManager(directory)
            runner = StressTestRunner(session)
            evidence = (EvidenceDefinition("workload", "Workload", "log", "workload", "synthetic"),)
            item = definition(evidence=evidence, program=sys.executable, arguments=("-c", "import time; time.sleep(30)"))
            runner.enqueue([item], environment_status=EnvironmentStatus.READY)
            runner.start_current()
            self.assertTrue(process_until(lambda: runner.evidence_manager.records[0].status == CollectorStatus.RUNNING))
            attempt_dir = runner.current_paths.attempt_dir
            runner.stop()
            self.assertTrue(process_until(lambda: item.runtime_status == RuntimeStatus.STOPPED))
            result = json.loads((attempt_dir / "result.json").read_text())
            self.assertEqual(result["status"], "STOPPED")

    def test_47_48_vr_and_vmo_placeholders(self):
        with tempfile.TemporaryDirectory() as directory:
            page = SystemStressPage(evidence_root=directory)
            self.assertIn("No VR Test Case Catalog loaded.", page.vr_page.findChildren(type(page.root_label))[1].text())
            labels = [label.text() for label in page.vmo_page.findChildren(type(page.root_label))]
            self.assertIn("No VMO Test Case Catalog loaded.", labels)
            page.close()

    def test_49_50_main_window_opens_stress_and_existing_pages_construct(self):
        window = MainWindow()
        index = len(window.NAV_ITEMS) - 1
        self.assertEqual(window.NAV_ITEMS[index][1], "Stress Test")
        self.assertEqual(window.nav_buttons[index].toolTip(), "System Stress Test")
        window.navigate_to(index)
        self.assertIs(window.pages.currentWidget(), window.system_stress_page)
        window.system_stress_page.navigate_requested.emit(0)
        self.assertIs(window.pages.currentWidget(), window.dashboard_page)
        window.navigate_to(index)
        self.assertIs(window.pages.currentWidget(), window.system_stress_page)
        self.assertTrue(hasattr(window, "camera_page"))
        self.assertTrue(hasattr(window, "lidar_page"))
        window.system_stress_page.shutdown()
        window.lidar_stream_service.shutdown()
        window.lidar_page.shutdown()
        window.jetson_service.shutdown()
        window.close()


if __name__ == "__main__":
    unittest.main()
