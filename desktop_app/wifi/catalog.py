"""Read the authoritative VD Wi-Fi suite; environments have independent catalogs."""
from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(__file__).resolve().parents[2] / "testcases/wifi/source/Jetson_WiFi_TCs_VD_completed.xlsx"
AUTO_SOURCE = Path(__file__).resolve().parents[2] / "testcases/wifi/source/wifi_vd_tcs.ods"
ENVIRONMENTS = ("VD", "VMO", "VR")
EXCLUDED_IDS = frozenset({"TC-WIFI-C08"})
# Mixed requirement-dependent reviews remain GUIDED. The remaining source
# acceptance criteria are objective and are not downgraded merely because they
# need multiple commands, parsing, comparison, iperf3, or profile changes.
EXISTING_GUIDED_IDS = frozenset({"TC-WIFI-C02", "TC-WIFI-C04"})
PLACEHOLDER = re.compile(r"^\s*<[^>]+>\s*$")
AUTO_SHEETS = {
    "JET_2.4G": ("AP_24G", "AP", "2.4G", 13, "TC-JET-24G-"),
    "JET_5G": ("AP_5G", "AP", "5G", 13, "TC-JET-5G-"),
    "JET_STA_2.4G": ("STA_24G", "STA", "2.4G", 8, "TC-JET-STA-24G-"),
    "JET_STA_5G": ("STA_5G", "STA", "5G", 8, "TC-JET-STA-5G-"),
}


@dataclass(frozen=True)
class WifiTestCase:
    test_id: str
    target: str
    category: str
    name: str
    purpose: str
    preconditions: str
    requirement: str
    procedure: str
    expected: str
    actual: str
    source_status: str
    notes: str
    mode: str
    phase: str
    source: str
    source_fields: dict[str, str]
    suite: str = "EXISTING"
    band: str = ""
    wifi_role: str = ""
    catalog_group: str = ""
    required_equipment: str = ""
    evidence_commands: str = ""
    traceability: str = ""

    @property
    def group(self) -> str:
        return self.category


def _rows(sheet):
    iterator = sheet.values
    headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
    for values in iterator:
        row = {key: str(value).strip() if value is not None else "" for key, value in zip(headers, values) if key}
        if any(row.values()):
            yield row


def load_config(source: Path = SOURCE) -> dict[str, str]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        result = {}
        for row in list(workbook["Config"].values)[1:]:
            if row[0]:
                value = str(row[1]).strip() if row[1] is not None else ""
                result[str(row[0]).strip()] = "NOT CONFIGURED" if not value or PLACEHOLDER.fullmatch(value) else value
        return result
    finally:
        workbook.close()


def load_execution_order(source: Path = SOURCE) -> list[dict[str, str]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return list(_rows(workbook["Execution Order"]))
    finally:
        workbook.close()


def source_rows(source: Path = SOURCE) -> list[dict[str, str]]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = list(_rows(workbook["Consolidated TCs"]))
        ids = [row["ID TC"] for row in rows]
        if len(ids) != len(set(ids)):
            raise ValueError("Duplicate Wi-Fi source test ID")
        return rows
    finally:
        workbook.close()


def _phase(test_id: str, order: list[dict[str, str]]) -> str:
    number = int(test_id.rsplit("C", 1)[1])
    for item in order:
        numbers = [int(value) for value in re.findall(r"C(\d+)", item["Consolidated TCs"])]
        if numbers and min(numbers) <= number <= max(numbers):
            return item["Phase"]
    return ""


def load_catalog(environment: str, source: Path = SOURCE) -> list[WifiTestCase]:
    if environment not in ENVIRONMENTS:
        raise ValueError(f"Unknown Wi-Fi environment: {environment}")
    if environment != "VD":
        return []
    order = load_execution_order(source)
    cases = []
    for row in source_rows(source):
        test_id = row["ID TC"]
        if test_id in EXCLUDED_IDS:
            continue
        number = int(test_id.rsplit("C", 1)[1])
        mode = "GUIDED" if test_id in EXISTING_GUIDED_IDS else "AUTO"
        cases.append(WifiTestCase(
            test_id=test_id, target="Jetson", category=row["Group"], name=row["Test Case"],
            purpose=row["Purpose"], preconditions=row["Preconditions"],
            requirement=row["Requirement"], procedure=row["Operation / Procedure"],
            expected=row["Expected Result"], actual=row["Actual Results"],
            source_status=row["Status"], notes=row["Notes"], mode=mode,
            phase=_phase(test_id, order), source=source.name, source_fields=row,
        ))
    return cases


_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"


def _ods_rows(table: ET.Element) -> list[list[str]]:
    """Read ODS rows without normalising or rewriting source cell text."""
    rows: list[list[str]] = []
    for row in table.findall(f"{{{_TABLE}}}table-row"):
        values: list[str] = []
        for cell in row.findall(f"{{{_TABLE}}}table-cell"):
            paragraphs = ["".join(item.itertext()) for item in cell.findall(f"{{{_TEXT}}}p")]
            value = "\n".join(paragraphs).strip()
            repeat = int(cell.attrib.get(f"{{{_TABLE}}}number-columns-repeated", "1"))
            values.extend([value] * repeat)
        if any(values):
            rows.append(values)
    return rows


def auto_source_rows(source: Path = AUTO_SOURCE) -> dict[str, list[dict[str, str]]]:
    """Load all required ODS sheets, failing on every structural/count mismatch."""
    if not source.is_file():
        raise FileNotFoundError(f"Wi-Fi AUTO source is missing: {source}")
    try:
        with zipfile.ZipFile(source) as archive:
            root = ET.fromstring(archive.read("content.xml"))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as error:
        raise ValueError(f"Invalid Wi-Fi AUTO ODS: {source}: {error}") from error
    tables = {table.attrib.get(f"{{{_TABLE}}}name", ""): table
              for table in root.findall(f".//{{{_TABLE}}}table")}
    missing = set(AUTO_SHEETS) - set(tables)
    if missing:
        raise ValueError(f"Wi-Fi AUTO ODS missing sheet(s): {', '.join(sorted(missing))}")
    unexpected = set(tables) - set(AUTO_SHEETS)
    if unexpected:
        raise ValueError(f"Wi-Fi AUTO ODS has unexpected sheet(s): {', '.join(sorted(unexpected))}")
    result: dict[str, list[dict[str, str]]] = {}
    for sheet_name, (_group, _role, _band, expected_count, prefix) in AUTO_SHEETS.items():
        rows = _ods_rows(tables[sheet_name])
        header_index = next((index for index, row in enumerate(rows) if row and row[0] == "ID TC"), None)
        if header_index is None:
            raise ValueError(f"Wi-Fi AUTO sheet {sheet_name} has no ID TC header")
        headers = rows[header_index]
        required = {"ID TC", "Host", "Band", "Group", "Category", "Test Case", "Purpose",
                    "Precondition", "Required Equipment", "Operation Procedure", "Expected Result",
                    "Operation Evidence / Commands", "Requirement Traceability", "Status"}
        if not required.issubset(headers):
            raise ValueError(f"Wi-Fi AUTO sheet {sheet_name} missing columns: {sorted(required - set(headers))}")
        imported: list[dict[str, str]] = []
        for source_row, values in enumerate(rows[header_index + 1:], header_index + 2):
            row = {name: values[index].strip() if index < len(values) else ""
                   for index, name in enumerate(headers) if name}
            if not row.get("ID TC"):
                raise ValueError(f"Wi-Fi AUTO sheet {sheet_name} row {source_row} has no ID TC")
            imported.append(row)
        expected_ids = [f"{prefix}{index:03d}" for index in range(1, expected_count + 1)]
        ids = [row["ID TC"] for row in imported]
        if len(imported) != expected_count or ids != expected_ids:
            raise ValueError(f"Wi-Fi AUTO sheet {sheet_name}: expected {expected_count} ordered IDs "
                             f"{expected_ids[0]}...{expected_ids[-1]}, imported {len(imported)}: {ids}")
        result[sheet_name] = imported
    total = sum(map(len, result.values()))
    if total != 42:
        raise ValueError(f"Wi-Fi AUTO ODS must import exactly 42 test cases; imported {total}")
    return result


def load_auto_catalog(source: Path = AUTO_SOURCE) -> list[WifiTestCase]:
    cases: list[WifiTestCase] = []
    for sheet_name, rows in auto_source_rows(source).items():
        catalog_group, role, band, _count, _prefix = AUTO_SHEETS[sheet_name]
        for row in rows:
            cases.append(WifiTestCase(
                test_id=row["ID TC"], target="Jetson", category=row["Category"],
                name=row["Test Case"], purpose=row["Purpose"], preconditions=row["Precondition"],
                requirement=row["Requirement Traceability"], procedure=row["Operation Procedure"],
                expected=row["Expected Result"], actual="", source_status=row["Status"], notes="",
                mode="AUTO", phase=row["Group"], source=source.name, source_fields=row,
                suite="AUTO", band=band, wifi_role=role, catalog_group=catalog_group,
                required_equipment=row["Required Equipment"],
                evidence_commands=row["Operation Evidence / Commands"],
                traceability=row["Requirement Traceability"],
            ))
    ids = [case.test_id for case in cases]
    if len(cases) != 42 or len(ids) != len(set(ids)):
        raise ValueError(f"Wi-Fi AUTO catalog integrity failure: {len(cases)} rows, {len(set(ids))} unique IDs")
    return cases
